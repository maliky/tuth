"""Synchronize administrative staff roles from TU directory sources."""

from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Mapping, TypeAlias

from django.contrib.auth.models import Group, User
from django.db import transaction
from openpyxl import load_workbook

from app.people.models import RoleAssignment, Staff
from app.people.utils import name_parts_from_row
from app.shared.auth.perms import UserRole

AuditRowT: TypeAlias = dict[str, str]
RoleLabelT: TypeAlias = str

VALID_EMAIL_RE = re.compile(r"^[^@\s,;]+@tubmanu\.edu\.lr$", re.IGNORECASE)
EMAIL_SPLIT_RE = re.compile(r"[,;\s]+")
ROLE_OFFICE_RE = re.compile(
    r"OFFICE OF THE REGISTRAR|OFFICE OF ENROLLMENT MANAGEMENT|"
    r"DIVISION OF ENROLLMENT MANAGEMENT|BUDGET AND FINANCE",
    re.IGNORECASE,
)
EMAIL_CORRECTIONS = {
    "derskin@tubmanu.edu.lr": "derskine@tubmanu.edu.lr",
    "jelloitt@tubmanu.edu.lr": "jelliott@tubmanu.edu.lr",
}
NAME_CORRECTIONS = {
    "Joe G. Elloitt": "Joe G. Elliott",
}
MANUAL_REVIEW_EMAILS = {
    "gfreeman@tubman.edu.lr": "conflicting Google and directory email evidence",
}


@dataclass(frozen=True)
class DirectoryAdminPerson:
    """Normalized administrative directory source row."""

    source: str
    row_id: str
    name: str
    email: str
    username: str
    position: str
    phone: str
    division: str
    role: RoleLabelT
    notes: str = ""


@dataclass(frozen=True)
class SyncOptions:
    """Runtime options for administrative directory role synchronization."""

    xlsx_path: Path
    google_csv_path: Path
    audit_path: Path
    apply: bool
    password: str
    today: date


@dataclass
class SyncState:
    """Mutable state shared across one administrative role sync run."""

    options: SyncOptions
    audit_rows: list[AuditRowT]
    counters: Counter[str]


def sync_administrative_directory_roles(options: SyncOptions) -> Counter[str]:
    """Synchronize registrar, finance, enrollment, and cashier staff roles."""
    state = SyncState(options=options, audit_rows=[], counters=Counter())
    google_emails = load_google_staff_emails(options.google_csv_path)
    source_people = collect_administrative_people(options.xlsx_path, google_emails)

    with transaction.atomic():
        for person in source_people:
            if not person.role:
                _audit(state, "manual_review", person, "role", person.notes)
                continue
            if not _valid_email(person.email):
                _audit(state, "manual_review", person, "email", person.notes)
                continue
            sync_person(state, person)
        if not options.apply:
            transaction.set_rollback(True)

    write_audit(options.audit_path, state.audit_rows)
    return state.counters


def load_google_staff_emails(path: Path) -> dict[str, str]:
    """Return Google staff email hints keyed by normalized full name."""
    emails: dict[str, str] = {}
    if not path.exists():
        return emails
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            org_unit = _text(row.get("Org Unit Path [Required]")).lower()
            if "staff" not in org_unit and "faculty" not in org_unit:
                continue
            first = _text(row.get("First Name [Required]"))
            last = _text(row.get("Last Name [Required]"))
            name = f"{first} {last}".strip()
            email = _text(row.get("Email Address [Required]")).lower()
            if name and _valid_email(email):
                emails[_name_key(name)] = email
    return emails


def collect_administrative_people(
    xlsx_path: Path,
    google_emails: Mapping[str, str],
) -> list[DirectoryAdminPerson]:
    """Return administrative staff candidates from the XLSX directory."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(max_row=1))]
    people: list[DirectoryAdminPerson] = []
    for row_no, cells in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        row = {headers[index]: cell for index, cell in enumerate(cells)}
        person = _person_from_xlsx_row(str(row_no), row, google_emails)
        if person is not None:
            people.append(person)
    return people


def sync_person(state: SyncState, person: DirectoryAdminPerson) -> None:
    """Create/update a user, staff profile, group membership, and role row."""
    user, created = _ensure_user(state, person)
    _audit(
        state,
        "created" if created else "updated",
        person,
        "user",
        user.username,
        object_id=user.pk,
    )
    staff = _ensure_staff(state, user, person)
    _ensure_groups(state, user, [UserRole.STAFF.value.label, person.role], person)
    _ensure_role(state, user, _group_for_role(person.role), person)
    _audit(state, "updated", person, "staff", staff.staff_id, object_id=staff.pk)


def write_audit(path: Path, rows: list[AuditRowT]) -> None:
    """Write stable CSV audit output for review."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "action",
        "status",
        "source",
        "row_id",
        "name",
        "email",
        "username",
        "role",
        "message",
        "object_id",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _person_from_xlsx_row(
    row_id: str,
    row: Mapping[str, object],
    google_emails: Mapping[str, str],
) -> DirectoryAdminPerson | None:
    raw_name = _text(row.get("Name"))
    position = _text(row.get("Position"))
    division = _text(row.get("Offices/Divisions"))
    if not ROLE_OFFICE_RE.search(division):
        return None
    if not raw_name or raw_name.lower() == "vacant":
        return _manual_person(row_id, raw_name or "Vacant", position, division, "vacant")

    name = NAME_CORRECTIONS.get(raw_name, raw_name)
    email, notes = _resolve_email(_text(row.get("Email")), name, google_emails)
    role, role_notes = _role_for(position, division, email)
    if role_notes:
        note_parts = dict.fromkeys(part for part in (notes, role_notes) if part)
        notes = "; ".join(note_parts)
    return DirectoryAdminPerson(
        source="xlsx",
        row_id=row_id,
        name=name,
        email=email,
        username=email.split("@", 1)[0] if _valid_email(email) else "",
        position=position,
        phone=_text(row.get("CellNo")),
        division=division,
        role=role,
        notes=notes,
    )


def _manual_person(
    row_id: str,
    name: str,
    position: str,
    division: str,
    reason: str,
) -> DirectoryAdminPerson:
    return DirectoryAdminPerson(
        source="xlsx",
        row_id=row_id,
        name=name,
        email="",
        username="",
        position=position,
        phone="",
        division=division,
        role="",
        notes=reason,
    )


def _resolve_email(
    raw_email: str,
    name: str,
    google_emails: Mapping[str, str],
) -> tuple[str, str]:
    raw_email = raw_email.lower()
    if raw_email in MANUAL_REVIEW_EMAILS:
        return raw_email, MANUAL_REVIEW_EMAILS[raw_email]
    corrected = EMAIL_CORRECTIONS.get(raw_email)
    if corrected:
        return corrected, f"corrected email from {raw_email}"

    emails = [value.lower() for value in EMAIL_SPLIT_RE.split(raw_email) if value]
    valid_emails = [email for email in emails if _valid_email(email)]
    if valid_emails:
        note = ""
        if len(valid_emails) > 1:
            note = f"multiple emails in source: {', '.join(valid_emails[1:])}"
        return valid_emails[0], note

    google_email = google_emails.get(_name_key(name))
    if google_email:
        return google_email, "corrected email from Google export"
    return raw_email, "invalid or missing email"


def _role_for(position: str, division: str, email: str) -> tuple[RoleLabelT, str]:
    position_upper = position.upper()
    division_upper = division.upper()
    if email in MANUAL_REVIEW_EMAILS:
        return "", MANUAL_REVIEW_EMAILS[email]
    if "OFFICE OF THE REGISTRAR" in division_upper:
        if position_upper.strip() == "REGISTRAR":
            return UserRole.REGISTRAR.value.label, ""
        if "REGISTRAR" in position_upper or "REGISTRATION" in position_upper:
            return UserRole.REGISTRAR_OFFICER.value.label, ""
        return "", "registrar office position is not an officer role"
    if "ENROLLMENT" in division_upper:
        if (
            "ASSOCIATE VICE PRESIDENT" in position_upper
            or "ASSISTANT VICE PRESIDENT" in position_upper
        ):
            return UserRole.ENROLLMENT.value.label, ""
        if any(
            token in position_upper
            for token in ("ADMISSION", "ENROLLEMENT", "ENROLLMENT", "STUDENT RECORDS")
        ):
            return UserRole.ENROLLMENT_OFFICER.value.label, ""
        return "", "enrollment office position is not an officer role"
    if "BUDGET AND FINANCE" in division_upper:
        if "CASHIER" in position_upper:
            return UserRole.CASHIER.value.label, ""
        if (
            "ASSOCIATE VICE PRESIDENT" in position_upper
            or "CHIEF ACCOUNTANT" in position_upper
        ):
            return UserRole.FINANCE.value.label, ""
        if "ACCOUNT" in position_upper:
            return UserRole.FINANCE_OFFICER.value.label, ""
        return "", "finance office position is not an officer role"
    return "", "office is not mapped to an administrative role"


def _ensure_user(
    state: SyncState,
    person: DirectoryAdminPerson,
) -> tuple[User, bool]:
    username = person.username or _username_from_name(person.name)
    user = User.objects.filter(username__iexact=username).first()
    if user is None and person.email:
        user = User.objects.filter(email__iexact=person.email).first()
    name = name_parts_from_row({}, raw_name=person.name).to_dict(full=False)
    defaults = {
        "username": username,
        "first_name": name["first_name"],
        "last_name": name["last_name"],
        "email": person.email,
        "is_staff": True,
        "is_active": True,
    }
    if user is None:
        user = User(**defaults)
        user.set_password(state.options.password)
        if state.options.apply:
            user.save()
        return user, True
    for field, value in defaults.items():
        setattr(user, field, value)
    if state.options.apply:
        user.save(update_fields=list(defaults))
    return user, False


def _ensure_staff(
    state: SyncState,
    user: User,
    person: DirectoryAdminPerson,
) -> Staff:
    staff = Staff.objects.filter(user=user).first() if user.pk else None
    if staff is None:
        staff = Staff(user=user)
    staff.position = _limit(person.position, 50)
    staff.division = _limit(person.division, 100)
    if person.phone:
        staff.phone_number = person.phone
    if state.options.apply:
        staff.save()
    return staff


def _ensure_groups(
    state: SyncState,
    user: User,
    group_names: Iterable[str],
    person: DirectoryAdminPerson,
) -> None:
    for group_name in dict.fromkeys(group_names):
        group, _created = Group.objects.get_or_create(name=group_name)
        if user.pk and user.groups.filter(pk=group.pk).exists():
            continue
        if state.options.apply and user.pk:
            user.groups.add(group)
        _audit(state, "role_added", person, "group", group.name, object_id=group.pk)


def _ensure_role(
    state: SyncState,
    user: User,
    group: Group,
    person: DirectoryAdminPerson,
) -> None:
    exists = bool(
        user.pk
        and RoleAssignment.objects.filter(
            user=user,
            group=group,
            college__isnull=True,
            department__isnull=True,
            end_date__isnull=True,
        ).exists()
    )
    if exists:
        return
    if state.options.apply:
        RoleAssignment.objects.create(
            user=user,
            group=group,
            start_date=state.options.today,
        )
    _audit(state, "role_added", person, group.name, "")


def _group_for_role(role: str) -> Group:
    for user_role in UserRole:
        if user_role.value.label == role:
            return user_role.value.group
    return Group.objects.get_or_create(name=role)[0]


def _audit(
    state: SyncState,
    action: str,
    person: DirectoryAdminPerson,
    role: str,
    message: str,
    *,
    object_id: int | None = None,
) -> None:
    state.counters[action] += 1
    state.audit_rows.append(
        {
            "action": action,
            "status": "applied" if state.options.apply else "dry_run",
            "source": person.source,
            "row_id": person.row_id,
            "name": person.name,
            "email": person.email,
            "username": person.username,
            "role": role,
            "message": message,
            "object_id": str(object_id or ""),
        }
    )


def _valid_email(email: str) -> bool:
    return bool(VALID_EMAIL_RE.match(email))


def _name_key(name: str) -> str:
    parsed = name_parts_from_row({}, raw_name=name)
    return " ".join([parsed.first, parsed.middle, parsed.last]).lower().replace(".", "")


def _username_from_name(name: str) -> str:
    parsed = name_parts_from_row({}, raw_name=name)
    return f"{parsed.first[:1]}{parsed.last}".lower().replace(" ", "")


def _text(value: object) -> str:
    return str(value or "").strip()


def _limit(value: str, max_length: int) -> str:
    return value[:max_length]


__all__ = [
    "SyncOptions",
    "collect_administrative_people",
    "load_google_staff_emails",
    "sync_administrative_directory_roles",
]
