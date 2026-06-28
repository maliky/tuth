"""Synchronize academic staff roles from TU directory sources."""

from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Mapping, TypeAlias

from django.contrib.auth.models import Group, User
from django.db import transaction
from openpyxl import load_workbook

from app.academics.models import College, Department
from app.people.models import Faculty, RoleAssignment, Staff
from app.people.utils import mk_username, name_parts_from_row
from app.shared.auth.perms import UserRole

AuditRowT: TypeAlias = dict[str, str]

FACULTY_POSITION_RE = re.compile(
    r"Professor|Dean|Director|Instructor|Lecturer|Chair", re.IGNORECASE
)
VALID_EMAIL_RE = re.compile(r"^[^@\s]+@tubmanu\.edu\.lr$", re.IGNORECASE)
COLLEGE_BY_DIVISION = {
    "COLLEGE OF AGRICULTURE AND FOOD SCIENCE": "CAFS",
    "COLLEGE OF AGRICULTURE AND FOOD SCIENCES": "CAFS",
    "COLLEGE OF ARTS AND SCIENCES": "CAS",
    "COLLEGE OF BUSINESS & ADMINISTRATION": "CBA",
    "COLLEGE OF BUSINESS ADMINISTRATION": "CBA",
    "COLLEGE OF EDUCATION": "EDRCE",
    "COLLEGE OF ENGINEERING AND TECHNOLOGY": "CET",
    "COLLEGE OF HEALTH SCIENCES": "CHS",
}
CURRENT_DEANS = (
    ("CAFS", "Larry C. Hwang, Ph.D.", "lhwang@tubmanu.edu.lr", "lhwang", "Dean"),
    ("CBA", "Tukundane Benson", "tbenson@tubmanu.edu.lr", "tbenson", "Dean"),
    ("EDRCE", "Dr. Jerry Castor BU-AO", "jbuao@tubmanu.edu.lr", "jbuao", "Dean"),
    ("CET", "Babatunde Fatai Ogunbayo", "bogunbayo@tubmanu.edu.lr", "bogunbayo", "Dean"),
    (
        "CHS",
        "Neonora Ade Payne-Barclay",
        "npayne-barclay@tubmanu.edu.lr",
        "npayne-barclay",
        "OIC Dean",
    ),
)
VPAA = (
    "Maria Fe Rebecca D. Gueta, Ed.D",
    "mgueta@tubmanu.edu.lr",
    "mgueta",
    "Vice President / Professor",
)
CAS_DEAN_VACANT = "CAS"
CHAIR_DEPARTMENTS = {
    "fdorley@tubmanu.edu.lr": ("ECON", "BFIN"),
    "ssaad@tubmanu.edu.lr": ("PADM", "BUSA"),
    "gwlateh@tubmanu.edu.lr": ("ECD", "PEDU"),
    "wmoulton@tubmanu.edu.lr": ("GCED",),
    "wdavis@tubmanu.edu.lr": ("EENG",),
    "awneufville@tubmanu.edu.lr": ("CENG",),
    "calbilimi@tubmanu.edu.lr": ("CSEN",),
    "dseyaker@tubmanu.edu.lr": ("PUBH",),
    "lcopeland@tubmanu.edu.lr": ("NURS",),
}
AMBIGUOUS_CHAIRS = {
    "omoigua@tubmanu.edu.lr": "Sports Science & Coaching has no deterministic department mapping.",
    "fnymah@tubmanu.edu.lr": "Math & Natural Sciences spans multiple departments.",
    "bakpah@tubmanu.edu.lr": "Languages & Literature spans multiple department codes.",
    "oedosa@tubmanu.edu.lr": "Mechanical Engineering has two chair rows in the directory.",
    "solabisi@tubmanu.edu.lr": "Mechanical Engineering has two chair rows in the directory.",
}


@dataclass(frozen=True)
class DirectoryPerson:
    """Normalized staff/faculty source row."""

    source: str
    row_id: str
    name: str
    email: str
    username: str
    position: str
    phone: str
    division: str
    college_code: str
    is_chair: bool = False
    is_dean: bool = False


@dataclass(frozen=True)
class SyncOptions:
    """Runtime options for academic directory role synchronization."""

    xlsx_path: Path
    google_csv_path: Path
    audit_path: Path
    apply: bool
    password: str
    today: date


@dataclass
class SyncState:
    """Mutable command state used while writing audit rows."""

    options: SyncOptions
    audit_rows: list[AuditRowT]
    counters: Counter[str]


def sync_academic_directory_roles(options: SyncOptions) -> Counter[str]:
    """Synchronize academic Staff, Faculty, Dean, Chair, and VPAA records."""
    state = SyncState(options=options, audit_rows=[], counters=Counter())
    google_people = load_google_people(options.google_csv_path)
    source_people = collect_source_people(options.xlsx_path, google_people)

    with transaction.atomic():
        for person in source_people:
            if _valid_email(person.email):
                sync_person(state, person, role="")
            else:
                _audit(state, "manual_review", person, "invalid_email", person.email)
        for person in current_dean_people(google_people):
            sync_person(state, person, role=UserRole.DEAN.value.label)
        sync_person(state, vpaa_person(google_people), role=UserRole.VPAA.value.label)
        retire_stale_roles(state)
        if not options.apply:
            transaction.set_rollback(True)

    write_audit(options.audit_path, state.audit_rows)
    return state.counters


def load_google_people(path: Path) -> dict[str, DirectoryPerson]:
    """Return Google users keyed by normalized full name."""
    people: dict[str, DirectoryPerson] = {}
    if not path.exists():
        return people
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for line_no, row in enumerate(csv.DictReader(handle), start=2):
            first = _text(row.get("First Name [Required]"))
            last = _text(row.get("Last Name [Required]"))
            name = f"{first} {last}".strip()
            email = _text(row.get("Email Address [Required]")).lower()
            if not name or not email:
                continue
            people[_name_key(name)] = DirectoryPerson(
                source="google_csv",
                row_id=str(line_no),
                name=name,
                email=email,
                username=email.split("@", 1)[0],
                position=_text(row.get("Employee Title")),
                phone=_text(row.get("Mobile Phone") or row.get("Work Phone")),
                division=_text(
                    row.get("Department") or row.get("Org Unit Path [Required]")
                ),
                college_code="",
            )
    return people


def google_person_for_name(
    name: str, google_people: Mapping[str, DirectoryPerson]
) -> DirectoryPerson | None:
    """Return one deterministic Google match by parsed last name and first initial."""
    direct = google_people.get(_name_key(name))
    if direct is not None:
        return direct
    parsed = name_parts_from_row({}, raw_name=name)
    if not parsed.first or not parsed.last:
        return None
    matches = [
        person
        for person in google_people.values()
        if _same_last_and_initial(parsed, name_parts_from_row({}, raw_name=person.name))
    ]
    return matches[0] if len(matches) == 1 else None


def collect_source_people(
    xlsx_path: Path, google_people: Mapping[str, DirectoryPerson]
) -> list[DirectoryPerson]:
    """Return valid college faculty-like people from the XLSX directory."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(max_row=1))]
    rows: list[DirectoryPerson] = []
    for row_no, cells in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        row = {headers[index]: cell for index, cell in enumerate(cells)}
        person = _person_from_xlsx_row(str(row_no), row, google_people)
        if person is not None:
            rows.append(person)
    return rows


def current_dean_people(
    google_people: Mapping[str, DirectoryPerson],
) -> list[DirectoryPerson]:
    """Return current dean override rows, independent of old directory vacancies."""
    return [
        _override_person(
            google_people,
            source="dean_override",
            college_code=college_code,
            name=name,
            email=email,
            username=username,
            position=position,
            is_dean=True,
        )
        for college_code, name, email, username, position in CURRENT_DEANS
    ]


def vpaa_person(google_people: Mapping[str, DirectoryPerson]) -> DirectoryPerson:
    """Return current VPAA override row."""
    name, email, username, position = VPAA
    return _override_person(
        google_people,
        source="vpaa_override",
        college_code="",
        name=name,
        email=email,
        username=username,
        position=position,
    )


def sync_person(state: SyncState, person: DirectoryPerson, *, role: str) -> None:
    """Create or update User, Staff, Faculty, groups, and scoped roles."""
    user, created = _ensure_user(state, person)
    if created:
        _audit(state, "created", person, "user", user.username, object_id=user.pk)
    else:
        _audit(state, "updated", person, "user", user.username, object_id=user.pk)
    staff = _ensure_staff(state, user, person)
    group_names = [UserRole.STAFF.value.label]
    if role != UserRole.VPAA.value.label:
        _ensure_faculty(state, staff, person)
        _set_staff_department(state, staff, person)
        group_names.append(UserRole.FACULTY.value.label)
        if person.is_chair:
            if _is_ambiguous_chair(person):
                _audit_ambiguous_chair(state, person)
            else:
                group_names.append(UserRole.CHAIR.value.label)
                _ensure_chair_roles(state, user, person)
        if role == UserRole.DEAN.value.label:
            group_names.append(UserRole.DEAN.value.label)
            _ensure_role(state, user, UserRole.DEAN.value.group, person.college_code, "")
    elif role:
        group_names.append(role)
        _ensure_role(state, user, UserRole.VPAA.value.group, "", "")
    _ensure_groups(state, user, group_names, person)


def retire_stale_roles(state: SyncState) -> None:
    """End placeholder and replaced academic leadership assignments."""
    dean_group = UserRole.DEAN.value.group
    chair_group = UserRole.CHAIR.value.group
    current_dean_usernames = {username for _c, _n, _e, username, _p in CURRENT_DEANS}
    for role in RoleAssignment.objects.filter(group=dean_group, end_date__isnull=True):
        college_code = role.college.code if role.college else ""
        should_retire = college_code == CAS_DEAN_VACANT or (
            college_code in {row[0] for row in CURRENT_DEANS}
            and role.user.username not in current_dean_usernames
        )
        if should_retire:
            _retire_role(state, role, "stale_dean")
    for role in RoleAssignment.objects.filter(
        group=chair_group, user__username__startswith="test_", end_date__isnull=True
    ):
        _retire_role(state, role, "placeholder_chair")
    ambiguous_emails = set(AMBIGUOUS_CHAIRS)
    for role in RoleAssignment.objects.filter(
        group=chair_group,
        user__email__in=ambiguous_emails,
        end_date__isnull=True,
    ):
        _retire_role(state, role, "ambiguous_chair_scope")


def write_audit(path: Path, rows: list[AuditRowT]) -> None:
    """Write stable CSV audit output."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "action",
        "status",
        "source",
        "row_id",
        "name",
        "email",
        "username",
        "role",
        "college",
        "department",
        "message",
        "object_id",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _person_from_xlsx_row(
    row_id: str,
    row: Mapping[str, object],
    google_people: Mapping[str, DirectoryPerson],
) -> DirectoryPerson | None:
    name = _text(row.get("Name"))
    if not name or name.lower().startswith("showing") or name.lower() == "vacant":
        return None
    position = _text(row.get("Position"))
    division = _text(row.get("Offices/Divisions"))
    college_code = COLLEGE_BY_DIVISION.get(division.upper(), "")
    if not college_code or not FACULTY_POSITION_RE.search(position):
        return None
    email = _text(row.get("Email")).lower()
    matched_google = google_person_for_name(name, google_people)
    if not _valid_email(email) and matched_google and _valid_email(matched_google.email):
        email = matched_google.email
    return DirectoryPerson(
        source="xlsx",
        row_id=row_id,
        name=name,
        email=email,
        username=email.split("@", 1)[0] if _valid_email(email) else "",
        position=position,
        phone=_text(row.get("CellNo")),
        division=division,
        college_code=college_code,
        is_chair="chair" in position.lower(),
        is_dean=bool(re.search(r"\bdean\b", position, re.IGNORECASE)),
    )


def _override_person(
    google_people: Mapping[str, DirectoryPerson],
    *,
    source: str,
    college_code: str,
    name: str,
    email: str,
    username: str,
    position: str,
    is_dean: bool = False,
) -> DirectoryPerson:
    matched_google = google_people.get(_name_key(name))
    phone = matched_google.phone if matched_google else ""
    return DirectoryPerson(
        source=source,
        row_id="override",
        name=name,
        email=email.lower(),
        username=username,
        position=position,
        phone=phone,
        division="",
        college_code=college_code,
        is_dean=is_dean,
    )


def _ensure_user(state: SyncState, person: DirectoryPerson) -> tuple[User, bool]:
    username = person.username or _username_from_name(person.name)
    user = User.objects.filter(username=username).first()
    if user is None and person.email:
        user = User.objects.filter(email__iexact=person.email).first()
    name = name_parts_from_row({}, raw_name=person.name)
    defaults = {
        "username": username,
        "first_name": name.first,
        "last_name": name.last,
        "email": person.email,
        "is_staff": True,
        "is_active": True,
    }
    if user is None:
        user = User(**defaults)
        user.set_password(state.options.password)
        if state.options.apply:
            user.save()
        return user, True
    for field, value in defaults.items():
        setattr(user, field, value)
    if state.options.apply:
        user.save(update_fields=list(defaults))
    return user, False


def _ensure_staff(state: SyncState, user: User, person: DirectoryPerson) -> Staff:
    staff = Staff.objects.filter(user=user).first() if user.pk else None
    created = staff is None
    if staff is None:
        staff = Staff(user=user)
    staff.position = _limit(person.position, 50)
    staff.division = _limit(person.division, 100)
    if person.phone:
        staff.phone_number = person.phone
    if state.options.apply:
        staff.save()
    _audit(
        state,
        "created" if created else "updated",
        person,
        "staff",
        staff.staff_id,
        object_id=staff.pk,
    )
    return staff


def _ensure_faculty(state: SyncState, staff: Staff, person: DirectoryPerson) -> Faculty:
    faculty = Faculty.objects.filter(staff_profile=staff).first() if staff.pk else None
    created = faculty is None
    if faculty is None:
        faculty = Faculty(staff_profile=staff)
    college = College.objects.get(code=person.college_code)
    faculty.college = college
    if state.options.apply:
        faculty.save()
    _audit(
        state,
        "created" if created else "updated",
        person,
        "faculty",
        college.code,
        object_id=faculty.pk,
    )
    return faculty


def _set_staff_department(
    state: SyncState, staff: Staff, person: DirectoryPerson
) -> None:
    departments = _mapped_departments(person)
    if len(departments) != 1:
        return
    department = departments[0]
    if staff.department_id == department.id:
        return
    staff.department = department
    if state.options.apply:
        staff.save(update_fields=["department"])


def _ensure_groups(
    state: SyncState, user: User, group_names: Iterable[str], person: DirectoryPerson
) -> None:
    for group_name in dict.fromkeys(group_names):
        group, _created = Group.objects.get_or_create(name=group_name)
        if user.pk and user.groups.filter(pk=group.pk).exists():
            continue
        if state.options.apply and user.pk:
            user.groups.add(group)
        _audit(state, "role_added", person, "group", group.name, object_id=group.pk)


def _ensure_chair_roles(state: SyncState, user: User, person: DirectoryPerson) -> None:
    departments = _mapped_departments(person)
    if not departments:
        _ensure_role(state, user, UserRole.CHAIR.value.group, person.college_code, "")
        return
    for department in departments:
        _ensure_role(
            state,
            user,
            UserRole.CHAIR.value.group,
            person.college_code,
            department.code,
        )


def _ensure_role(
    state: SyncState,
    user: User,
    group: Group,
    college_code: str,
    department_code: str,
) -> None:
    college = College.objects.filter(code=college_code).first() if college_code else None
    department = (
        Department.objects.filter(college=college, code=department_code).first()
        if college and department_code
        else None
    )
    if department_code and department is None:
        fake = DirectoryPerson(
            "mapping",
            "",
            user.get_full_name(),
            user.email,
            user.username,
            "",
            "",
            "",
            college_code,
        )
        _audit(state, "manual_review", fake, "missing_department", department_code)
        return
    exists = bool(
        user.pk
        and RoleAssignment.objects.filter(
            user=user,
            group=group,
            college=college,
            department=department,
            end_date__isnull=True,
        ).exists()
    )
    if exists:
        return
    if state.options.apply:
        RoleAssignment.objects.create(
            user=user,
            group=group,
            college=college,
            department=department,
            start_date=state.options.today,
        )
    _audit(
        state,
        "role_added",
        DirectoryPerson(
            "role",
            "",
            user.get_full_name(),
            user.email,
            user.username,
            "",
            "",
            "",
            college_code,
        ),
        group.name,
        "",
        department_code,
    )


def _retire_role(state: SyncState, role: RoleAssignment, reason: str) -> None:
    role.end_date = state.options.today
    if state.options.apply:
        role.save(update_fields=["end_date"])
        if not RoleAssignment.objects.filter(
            user=role.user, group=role.group, end_date__isnull=True
        ).exists():
            role.user.groups.remove(role.group)
    college_code = role.college.code if role.college else ""
    dept_code = role.department.code if role.department else ""
    person = DirectoryPerson(
        "db",
        str(role.pk),
        role.user.get_full_name(),
        role.user.email,
        role.user.username,
        "",
        "",
        "",
        college_code,
    )
    _audit(
        state,
        "role_retired",
        person,
        role.group.name,
        reason,
        dept_code,
        object_id=role.pk,
    )


def _mapped_departments(person: DirectoryPerson) -> list[Department]:
    codes = CHAIR_DEPARTMENTS.get(person.email.lower(), ())
    if not codes:
        return []
    return list(
        Department.objects.filter(college__code=person.college_code, code__in=codes)
        .select_related("college")
        .order_by("code")
    )


def _audit_ambiguous_chair(state: SyncState, person: DirectoryPerson) -> None:
    message = AMBIGUOUS_CHAIRS.get(person.email.lower())
    if message:
        _audit(state, "manual_review", person, "chair_scope", message)


def _is_ambiguous_chair(person: DirectoryPerson) -> bool:
    """Return whether a chair row needs manual scope review before activation."""
    return person.email.lower() in AMBIGUOUS_CHAIRS


def _audit(
    state: SyncState,
    action: str,
    person: DirectoryPerson,
    role: str,
    message: str,
    department: str = "",
    *,
    object_id: int | None = None,
) -> None:
    state.counters[action] += 1
    state.audit_rows.append(
        {
            "action": action,
            "status": "applied" if state.options.apply else "dry_run",
            "source": person.source,
            "row_id": person.row_id,
            "name": person.name,
            "email": person.email,
            "username": person.username,
            "role": role,
            "college": person.college_code,
            "department": department,
            "message": message,
            "object_id": str(object_id or ""),
        }
    )


def _text(value: object) -> str:
    return str(value or "").strip()


def _valid_email(email: str) -> bool:
    return bool(VALID_EMAIL_RE.match(email))


def _name_key(name: str) -> str:
    parsed = name_parts_from_row({}, raw_name=name)
    return " ".join([parsed.first, parsed.middle, parsed.last]).lower().replace(".", "")


def _same_last_and_initial(left: object, right: object) -> bool:
    left_first = getattr(left, "first", "")
    right_first = getattr(right, "first", "")
    return (
        bool(left_first)
        and bool(right_first)
        and left_first[:1].lower() == right_first[:1].lower()
        and getattr(left, "last", "").lower() == getattr(right, "last", "").lower()
    )


def _username_from_name(name: str) -> str:
    parsed = name_parts_from_row({}, raw_name=name)
    return mk_username(parsed.first, parsed.last, parsed.middle, unique=True)


def _limit(value: str, max_length: int) -> str:
    return value[:max_length]


__all__ = [
    "SyncOptions",
    "collect_source_people",
    "load_google_people",
    "sync_academic_directory_roles",
]
